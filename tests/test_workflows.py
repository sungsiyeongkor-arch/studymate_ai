import io
from datetime import timedelta
import pytest
from pypdf import PdfWriter
from conftest import Client
from studymate.extensions import db
from studymate.models import (
    Job,
    Material,
    Segment,
    Artifact,
    Citation,
    Card,
    ReviewState,
    Attempt,
    now,
)
from studymate.services import jobs, learning
from studymate.services.provider import ProviderError


def extract(app):
    with app.app_context():
        job = jobs.claim()
        assert job is not None
        jobs.execute(job)
        assert job.status == "completed", job.error


def test_csrf_auth_tenant_isolation(app, client):
    assert app.test_client().get("/api/dashboard").status_code == 401
    assert client.raw.post("/api/courses", json={"title": "bad"}).status_code == 403
    lecture = client.setup_lecture()
    _, response = client.upload(lecture)
    material = response.json
    other = Client(app, "other@example.test")
    for path in [
        "/lectures/" + lecture["id"],
        "/materials/" + material["id"],
        "/materials/" + material["id"] + "/file",
    ]:
        assert other.send("GET", path).status_code == 404
    assert (
        other.send(
            "POST", "/lectures", {"course_id": lecture["course_id"], "title": "bad"}
        ).status_code
        == 404
    )


def test_upload_resume_extract_version_and_delete(app, client):
    lecture = client.setup_lecture()
    upload, response = client.upload(lecture)
    assert response.status_code == 201
    mid = response.json["id"]
    assert (
        client.send("POST", "/uploads/" + upload["id"] + "/complete").json["id"] == mid
    )
    extract(app)
    segment = client.send("GET", "/materials/" + mid).json["segments"][0]
    sid = segment["id"]
    result = client.send(
        "PATCH",
        "/segments/" + sid,
        {"text": "Corrected reading", "revision": 1, "reason": "transcription"},
    )
    assert result.status_code == 200
    assert len(result.json["revisions"]) == 2
    assert (
        client.send(
            "PATCH", "/segments/" + sid, {"text": "overwrite", "revision": 1}
        ).status_code
        == 409
    )
    assert client.send("GET", "/search?q=Corrected").json["items"][0]["id"] == sid
    assert client.send("GET", "/search?q=%25%25").json["items"] == []
    assert client.send("DELETE", "/materials/" + mid).status_code == 200
    assert client.send("GET", "/materials/" + mid + "/file").status_code == 404


def test_bad_pdf_and_size_rejected(client):
    lecture = client.setup_lecture()
    _, r = client.upload(lecture, b"not a pdf", "wrong.pdf")
    assert r.status_code == 400
    assert (
        client.send(
            "POST",
            "/uploads",
            {"lecture_id": lecture["id"], "name": "big.pdf", "size": 600 * 1024**2},
        ).status_code
        == 400
    )


def test_pdf_crop_render_and_personalization(app, client):
    lecture = client.setup_lecture()
    pdf = PdfWriter()
    pdf.add_blank_page(width=600, height=800)
    buffer = io.BytesIO()
    pdf.write(buffer)
    _, r = client.upload(lecture, buffer.getvalue(), "handwriting.pdf")
    mid = r.json["id"]
    extract(app)
    assert (
        client.send("GET", "/materials/" + mid + "/pages/1").content_type == "image/png"
    )
    assert client.send("GET", "/materials/" + mid + "/pages/2").status_code == 400
    assert (
        client.send(
            "POST",
            "/materials/" + mid + "/crop",
            {"page": 1, "bbox": [0.9, 0.9, 0.5, 0.5]},
        ).status_code
        == 400
    )
    r = client.send(
        "POST",
        "/materials/" + mid + "/crop",
        {"page": 1, "bbox": [0.1, 0.1, 0.5, 0.15]},
    )
    assert r.status_code == 201
    sid = r.json["id"]
    assert client.send("GET", "/segments/" + sid + "/image").status_code == 200
    data = {
        "text": "my handwriting",
        "revision": 1,
        "save_example": True,
        "own_handwriting": True,
    }
    assert client.send("PATCH", "/segments/" + sid, data).status_code == 400
    assert client.send("GET", "/segments/" + sid).json["segment"]["revision"] == 1
    client.send("PATCH", "/settings", {"personalization": True})
    assert client.send("PATCH", "/segments/" + sid, data).status_code == 200
    assert len(client.send("GET", "/settings").json["examples"]) == 1
    assert client.send("POST", "/segments/" + sid + "/recognize").status_code == 503
    client.send("PATCH", "/segments/" + sid, {"text": "new", "revision": 2})
    assert client.send("GET", "/settings").json["examples"] == []


def test_review_idempotency_conflict_hint_and_report(app, client):
    lecture = client.setup_lecture()
    c = client.send(
        "POST",
        "/cards",
        {
            "lecture_id": lecture["id"],
            "question": "What is recall?",
            "answer": "Retrieving information.",
        },
    ).json
    state = client.send("GET", "/reviews/due").json["items"][0]["state"]
    data = {
        "grade": 3,
        "confidence": 2,
        "answer": "Retrieval",
        "state_version": state["version"],
        "request_key": "same",
        "hint_used": True,
    }
    r = client.send("POST", "/cards/" + c["id"] + "/review", data)
    assert r.status_code == 201
    assert (
        client.send("POST", "/cards/" + c["id"] + "/review", data).json["id"]
        == r.json["id"]
    )
    assert (
        client.send(
            "POST",
            "/cards/" + c["id"] + "/review",
            {**data, "request_key": "different"},
        ).status_code
        == 409
    )
    with app.app_context():
        assert Attempt.query.count() == 1
        assert ReviewState.query.one().interval <= 7
    assert client.send("POST", "/cards/" + c["id"] + "/report").status_code == 200
    assert client.send("GET", "/reviews/due").json["items"] == []
    assert client.send("GET", "/dashboard").json["counts"]["attempts"] == 0


def test_missing_key_blocks_job_honestly(app, client):
    l = client.setup_lecture()
    client.upload(l)
    extract(app)
    client.send("POST", "/lectures/" + l["id"] + "/generate", {"kind": "summary"})
    with app.app_context():
        job = jobs.claim()
        jobs.execute(job)
        assert job.status == "blocked"
        assert Artifact.query.count() == 0


def test_artifact_generation_atomic_and_stale(app, client, monkeypatch):
    l = client.setup_lecture()
    _, r = client.upload(l)
    extract(app)
    sid = client.send("GET", "/materials/" + r.json["id"]).json["segments"][0]["id"]
    monkeypatch.setattr(jobs, "credentials", lambda _: "test")
    monkeypatch.setattr(
        learning,
        "call_json",
        lambda *a, **k: {
            "title": "Memory quiz",
            "questions": [
                {
                    "question": "Q",
                    "answer": "A",
                    "options": ["A", "B", "C", "D"],
                    "kind": "multiple_choice",
                    "source_ids": [sid],
                }
            ],
        },
    )
    client.send("POST", "/lectures/" + l["id"] + "/generate", {"kind": "quiz"})
    extract(app)
    with app.app_context():
        assert Artifact.query.count() == 1
        assert Citation.query.one().revision == 1
    assert len(client.send("GET", "/reviews/due").json["items"]) == 1
    client.send("PATCH", "/segments/" + sid, {"text": "updated", "revision": 1})
    assert client.send("GET", "/reviews/due").json["items"] == []
    assert client.send("GET", "/dashboard").json["counts"]["due"] == 0


def test_worker_failure_rolls_back_partial_artifacts(app, client, monkeypatch):
    l = client.setup_lecture()
    client.upload(l)
    extract(app)
    monkeypatch.setattr(jobs, "credentials", lambda _: "test")

    def partial(user_id, lecture_id, *args):
        db.session.add(
            Artifact(
                user_id=user_id,
                lecture_id=lecture_id,
                kind="summary",
                title="partial",
                content={},
            )
        )
        db.session.flush()
        raise ValueError("simulated failure")

    monkeypatch.setattr(jobs, "generate_batch", partial)
    client.send("POST", "/lectures/" + l["id"] + "/generate", {"kind": "summary"})
    with app.app_context():
        job = jobs.claim()
        jobs.execute(job)
        assert job.status == "failed"
        assert Artifact.query.count() == 0
        assert job.payload.get("completed_batches") == []


def test_lease_reclaim_and_fencing(app, client):
    l = client.setup_lecture()
    client.upload(l)
    with app.app_context():
        old = jobs.claim()
        token = old.lease_token
        assert jobs.claim() is None
        Job.query.filter_by(id=old.id).update(
            {"heartbeat": now() - timedelta(minutes=20)}
        )
        db.session.commit()
        replacement = jobs.claim()
        assert replacement.lease_token != token
        replacement._claimed_token = token
        with pytest.raises(jobs.LostLease):
            jobs.checkpoint(replacement, 1, 1)


def test_export_scoped_and_secret_free(app, client):
    client.setup_lecture()
    other = Client(app, "other@example.test")
    other.setup_lecture()
    r = client.send("GET", "/account/export")
    assert r.status_code == 200
    assert len(r.json["sm_courses"]) == 1
    assert "password_hash" not in r.get_data(as_text=True)


@pytest.mark.parametrize("kind", ["pptx", "docx"])
def test_office_document_extraction_and_source_type(app, client, kind):
    buffer = io.BytesIO()
    if kind == "pptx":
        from pptx import Presentation

        doc = Presentation()
        slide = doc.slides.add_slide(doc.slide_layouts[1])
        slide.shapes.title.text = "Working memory"
        slide.placeholders[1].text = "Practice retrieval."
    else:
        from docx import Document

        doc = Document()
        doc.add_paragraph("Working memory: Practice retrieval.")
    doc.save(buffer)
    lecture = client.setup_lecture()
    _, response = client.upload(lecture, buffer.getvalue(), "source." + kind)
    assert response.status_code == 201
    extract(app)
    source = client.send("GET", "/materials/" + response.json["id"]).json["segments"][0]
    detail = client.send("GET", "/segments/" + source["id"]).json
    assert detail["material_kind"] == kind
    assert "Working memory" in detail["segment"]["text"]
