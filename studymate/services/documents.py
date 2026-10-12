"""Bounded extraction. Unreadable pages stay visible as review-required records."""

import io
import zipfile
import threading
from pathlib import Path
from pypdf import PdfReader
from PIL import Image
from .storage import path

ALLOWED = {
    "pdf",
    "pptx",
    "docx",
    "txt",
    "md",
    "wav",
    "mp3",
    "m4a",
    "mp4",
    "webm",
    "ogg",
    "flac",
}
AUDIO = {"wav", "mp3", "m4a", "mp4", "webm", "ogg", "flac"}
_pdf_lock = threading.RLock()


def validate(file_path, name):
    ext = Path(name).suffix.lower().lstrip(".")
    if ext not in ALLOWED:
        raise ValueError(
            "PDF, PPTX, DOCX, TXT 또는 지원하는 녹음 파일을 선택해 주세요."
        )
    with open(file_path, "rb") as stream:
        header = stream.read(32)
    if ext == "pdf":
        if not header.startswith(b"%PDF-"):
            raise ValueError("유효한 PDF 파일이 아닙니다.")
        reader = PdfReader(file_path)
        if reader.is_encrypted:
            raise ValueError("암호를 해제한 PDF를 업로드해 주세요.")
        if len(reader.pages) > 1000:
            raise ValueError("PDF는 1,000쪽 이하로 나누어 주세요.")
    elif ext in {"pptx", "docx"}:
        with zipfile.ZipFile(file_path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 256 * 1024 * 1024:
                raise ValueError("압축 해제 크기가 너무 큽니다.")
            expected = "ppt/presentation.xml" if ext == "pptx" else "word/document.xml"
            if expected not in archive.namelist():
                raise ValueError("유효한 문서 형식이 아닙니다.")
    elif ext in {"txt", "md"}:
        if Path(file_path).stat().st_size > 10 * 1024 * 1024:
            raise ValueError("텍스트 파일은 10MB 이하로 올려 주세요.")
        Path(file_path).read_text(encoding="utf-8-sig")
    elif not (
        header.startswith((b"RIFF", b"ID3", b"OggS", b"fLaC", b"\x1aE\xdf\xa3"))
        or header[4:8] == b"ftyp"
        or (len(header) > 1 and header[0] == 255 and header[1] & 224 == 224)
    ):
        raise ValueError("지원하는 음성 형식인지 확인해 주세요.")
    return "audio" if ext in AUDIO else ext


def extract(material):
    source = path(material.storage_key)
    if material.kind == "pdf":
        reader = PdfReader(source)
        for i, page in enumerate(reader.pages, 1):
            yield {"page": i, "text": page.extract_text() or ""}
    elif material.kind == "pptx":
        from pptx import Presentation

        deck = Presentation(source)
        for i, slide in enumerate(deck.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(
                        " | ".join(cell.text for cell in row.cells)
                        for row in shape.table.rows
                    )
            yield {"page": i, "text": "\n".join(texts)}
    elif material.kind == "docx":
        from docx import Document

        doc = Document(source)
        text = "\n".join(p.text for p in doc.paragraphs)
        for table in doc.tables:
            text += "\n" + "\n".join(
                " | ".join(c.text for c in r.cells) for r in table.rows
            )
        for i in range(0, max(len(text), 1), 8000):
            yield {"page": None, "text": text[i : i + 8000]}
    else:
        text = source.read_text(encoding="utf-8-sig")
        for i in range(0, max(len(text), 1), 8000):
            yield {"page": None, "text": text[i : i + 8000]}


def render_pdf(material, page_number, bbox=None, scale=2):
    # PDFium calls must not execute concurrently within the same process.
    with _pdf_lock:
        return _render_pdf(material, page_number, bbox, scale)


def _render_pdf(material, page_number, bbox=None, scale=2):
    import pypdfium2 as pdfium

    doc = pdfium.PdfDocument(str(path(material.storage_key)))
    try:
        if not 1 <= page_number <= len(doc):
            raise ValueError("페이지 번호를 확인해 주세요.")
        page = doc[page_number - 1]
        try:
            width, height = page.get_size()
            crop = (0, 0, 0, 0)
            if bbox:
                x, y, w, h = bbox
                if not all(isinstance(v, (int, float)) for v in bbox) or not (
                    0 <= x < 1
                    and 0 <= y < 1
                    and w > 0
                    and h > 0
                    and x + w <= 1.00001
                    and y + h <= 1.00001
                ):
                    raise ValueError("선택 영역은 페이지 안에 있어야 합니다.")
                crop = (
                    x * width,
                    (1 - y - h) * height,
                    (1 - x - w) * width,
                    y * height,
                )
            target_w = width * (bbox[2] if bbox else 1)
            target_h = height * (bbox[3] if bbox else 1)
            scale = min(scale, 3000 / max(target_w, target_h))
            bitmap = page.render(scale=scale, crop=crop)
            try:
                image = bitmap.to_pil().convert("RGB")
                buffer = io.BytesIO()
                image.save(buffer, format="PNG")
                return buffer.getvalue()
            finally:
                bitmap.close()
        finally:
            page.close()
    finally:
        doc.close()
